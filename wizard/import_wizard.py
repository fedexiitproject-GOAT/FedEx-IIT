import base64
import csv
import io
import json
from datetime import datetime

from odoo import api, fields, models
from odoo.exceptions import UserError, ValidationError

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = None
    load_workbook = None


ENTITY_MODELS = {
    "students": "fedex.student",
    "faculty": "fedex.faculty",
    "staff": "fedex.staff",
    "projects": "fedex.project",
    "project_members": "fedex.project.member",
}

TEMPLATES = {
    "students": ["student_code", "name", "email", "phone", "degree", "institution", "department", "year_of_study", "grading_scale", "cgpa", "status", "technical_skills", "soft_skills", "domain_skills"],
    "faculty": ["faculty_code", "name", "email", "phone", "institution", "department", "categories"],
    "staff": ["staff_code", "name", "email", "phone", "designation", "department"],
    "projects": ["code", "name", "description", "centre", "type", "categories", "sponsor", "supervisor_email", "state", "priority", "start_date", "expected_end_date", "actual_end_date", "progress_manual"],
    "project_members": ["project_code", "person_email", "role", "start_date", "expected_end_date", "actual_end_date", "state", "certificate_issued"],
}

LOOKUP_MODELS = {
    "degree": "fedex.degree", "institution": "fedex.institution", "department": "fedex.department", "centre": "fedex.centre",
    "type": "fedex.project.type", "sponsor": "fedex.sponsor", "category": "fedex.project.category", "role": "fedex.project.role",
}


class FedexImportWizard(models.TransientModel):
    _name = "fedex.import.wizard"
    _description = "FedEx Import Wizard"

    entity = fields.Selection([(k, k.replace("_", " ").title()) for k in ENTITY_MODELS], required=True)
    file = fields.Binary(required=True, attachment=False)
    filename = fields.Char(required=True)
    create_missing_lookups = fields.Boolean(string="Create missing lookup values")
    import_mode = fields.Selection([("skip", "Skip duplicates"), ("update", "Update existing")], default="skip", required=True)
    preview_line_ids = fields.One2many("fedex.import.preview", "wizard_id")
    state = fields.Selection([("draft", "Upload"), ("preview", "Preview")], default="draft")
    error_report = fields.Binary(readonly=True)
    error_report_filename = fields.Char(readonly=True)

    def _require_xlsx(self):
        if load_workbook is None:
            raise UserError("XLSX import requires the openpyxl Python package. Install it before using XLSX import.")

    def _rows(self):
        self.ensure_one()
        raw = base64.b64decode(self.file or b"")
        name = (self.filename or "").lower()
        if name.endswith(".csv"):
            return list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
        if name.endswith(".xlsx"):
            self._require_xlsx()
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            values = list(sheet.values)
            if not values:
                return []
            headers = [str(v).strip() if v is not None else "" for v in values[0]]
            return [dict(zip(headers, row)) for row in values[1:] if any(v is not None and str(v).strip() for v in row)]
        raise UserError("Only CSV and XLSX files are supported.")

    @staticmethod
    def _norm(value):
        return str(value or "").strip()

    def _lookup(self, model_name, value, field="name", create=False, extra=None):
        value = self._norm(value)
        if not value:
            return self.env[model_name].browse()
        domain = [(field, "=ilike", value)]
        if extra:
            domain += extra
        rec = self.env[model_name].search(domain, limit=1)
        if rec:
            return rec
        if create:
            vals = {field: value}
            vals.update(extra and {k[0]: k[2] for k in extra if len(k) == 3 and k[1] == "="} or {})
            return self.env[model_name].create(vals)
        raise ValidationError("Unknown %s: %s" % (model_name.replace("fedex.", ""), value))

    def _split_names(self, value):
        return [x.strip() for x in self._norm(value).split(",") if x.strip()]

    def _date(self, value):
        value = self._norm(value)
        if not value:
            return False
        if isinstance(value, datetime):
            return value.date()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                continue
        raise ValidationError("Invalid date: %s. Use YYYY-MM-DD." % value)

    def _find_duplicate(self, row):
        model = self.env[ENTITY_MODELS[self.entity]]
        if self.entity in ("students", "faculty", "staff"):
            email = self._norm(row.get("email"))
            code_field = {"students": "student_code", "faculty": "faculty_code", "staff": "staff_code"}[self.entity]
            code = self._norm(row.get(code_field))
            if email:
                rec = model.search([("email", "=ilike", email)], limit=1)
                if rec:
                    return rec
            if code:
                return model.search([(code_field, "=", code)], limit=1)
        elif self.entity == "projects":
            code = self._norm(row.get("code"))
            return model.search([("code", "=", code)], limit=1) if code else model.browse()
        return model.browse()

    def _build_vals(self, row):
        e = self.entity
        name = self._norm(row.get("name"))
        if e == "students":
            institution = self._lookup("fedex.institution", row.get("institution"), create=self.create_missing_lookups)
            department = self._lookup("fedex.department", row.get("department"), create=self.create_missing_lookups, extra=[("institution_id", "=", institution.id)]) if self._norm(row.get("department")) else self.env["fedex.department"].browse()
            vals = {"name": name, "email": self._norm(row.get("email")), "phone": self._norm(row.get("phone")), "degree_id": self._lookup("fedex.degree", row.get("degree"), create=self.create_missing_lookups).id, "institution_id": institution.id, "department_id": department.id or False, "year_of_study": int(row.get("year_of_study")) if self._norm(row.get("year_of_study")) else False, "grading_scale": self._norm(row.get("grading_scale")) or "cgpa10", "cgpa": float(row.get("cgpa")) if self._norm(row.get("cgpa")) else 0, "status": self._norm(row.get("status")) or "applicant"}
            for field, skill_type in (("technical_skill_ids", "technical"), ("soft_skill_ids", "soft"), ("domain_ids", "domain")):
                commands = []
                for skill in self._split_names(row.get(field.replace("_ids", "_skills"))):
                    rec = self._lookup("fedex.skill", skill, create=self.create_missing_lookups, extra=[("skill_type", "=", skill_type)])
                    commands.append(rec.id)
                vals[field] = [(6, 0, commands)]
            if self._norm(row.get("student_code")):
                vals["student_code"] = self._norm(row.get("student_code"))
            return vals
        if e == "faculty":
            institution = self._lookup("fedex.institution", row.get("institution"), create=self.create_missing_lookups)
            department = self._lookup("fedex.department", row.get("department"), create=self.create_missing_lookups, extra=[("institution_id", "=", institution.id)]) if self._norm(row.get("department")) else self.env["fedex.department"].browse()
            vals = {"name": name, "email": self._norm(row.get("email")), "phone": self._norm(row.get("phone")), "institution_id": institution.id, "department_id": department.id or False}
            cats = []
            for category in self._split_names(row.get("categories")):
                cats.append(self._lookup("fedex.project.category", category, create=self.create_missing_lookups).id)
            vals["category_ids"] = [(6, 0, cats)]
            if self._norm(row.get("faculty_code")): vals["faculty_code"] = self._norm(row.get("faculty_code"))
            return vals
        if e == "staff":
            department = self._lookup("fedex.department", row.get("department"), create=self.create_missing_lookups) if self._norm(row.get("department")) else self.env["fedex.department"].browse()
            vals = {"name": name, "email": self._norm(row.get("email")), "phone": self._norm(row.get("phone")), "designation": self._norm(row.get("designation")), "department_id": department.id or False}
            if self._norm(row.get("staff_code")): vals["staff_code"] = self._norm(row.get("staff_code"))
            return vals
        if e == "projects":
            type_rec = self._lookup("fedex.project.type", row.get("type"), create=self.create_missing_lookups)
            centre = self._lookup("fedex.centre", row.get("centre"), create=self.create_missing_lookups)
            vals = {"name": name, "description": self._norm(row.get("description")), "centre_id": centre.id, "type_id": type_rec.id, "state": self._norm(row.get("state")) or "yet_to_start", "priority": self._norm(row.get("priority")) or "normal", "start_date": self._date(row.get("start_date")), "expected_end_date": self._date(row.get("expected_end_date")), "actual_end_date": self._date(row.get("actual_end_date")), "progress_manual": float(row.get("progress_manual")) if self._norm(row.get("progress_manual")) else 0}
            cats = [self._lookup("fedex.project.category", c, create=self.create_missing_lookups).id for c in self._split_names(row.get("categories"))]
            vals["category_ids"] = [(6, 0, cats)]
            sponsor = self._norm(row.get("sponsor"))
            if sponsor: vals["sponsor_id"] = self._lookup("fedex.sponsor", sponsor, create=self.create_missing_lookups).id
            supervisor_email = self._norm(row.get("supervisor_email"))
            if supervisor_email:
                faculty = self.env["fedex.faculty"].search([("email", "=ilike", supervisor_email)], limit=1)
                if not faculty: raise ValidationError("Unknown faculty email: %s" % supervisor_email)
                vals["supervisor_id"] = faculty.id
            code = self._norm(row.get("code"))
            if code: vals["code"] = code
            return vals
        if e == "project_members":
            project = self.env["fedex.project"].search([("code", "=", self._norm(row.get("project_code")))], limit=1)
            if not project: raise ValidationError("Unknown project code: %s" % row.get("project_code"))
            partner = self.env["res.partner"].search([("email", "=ilike", self._norm(row.get("person_email")))], limit=1)
            if not partner: raise ValidationError("Unknown person email: %s" % row.get("person_email"))
            role = self._lookup("fedex.project.role", row.get("role"), create=self.create_missing_lookups)
            return {"project_id": project.id, "partner_id": partner.id, "role_id": role.id, "start_date": self._date(row.get("start_date")) or fields.Date.context_today(self), "expected_end_date": self._date(row.get("expected_end_date")), "actual_end_date": self._date(row.get("actual_end_date")), "state": self._norm(row.get("state")) or "active", "certificate_issued": str(row.get("certificate_issued") or "").lower() in ("1", "true", "yes")}
        raise ValidationError("Unsupported import entity.")

    def action_validate(self):
        self.ensure_one()
        self.preview_line_ids.unlink()
        rows = self._rows()
        if not rows:
            raise UserError("The uploaded file contains no data rows.")
        headers = set(rows[0].keys())
        required = set(TEMPLATES[self.entity])
        missing = required - headers
        if missing:
            raise UserError("Missing required template columns: %s" % ", ".join(sorted(missing)))
        seen = set()
        preview = []
        errors = []
        for row_number, row in enumerate(rows, 2):
            duplicate = self._find_duplicate(row)
            key = json.dumps(row, sort_keys=True, default=str)
            status = "update" if duplicate else "new"
            reason = "Existing record will be updated." if duplicate else ""
            if key in seen:
                status, reason = "duplicate", "Duplicate row in uploaded file."
            seen.add(key)
            try:
                self._build_vals(row)
            except Exception as exc:
                status, reason = "error", str(exc)
                errors.append({"row": row_number, "reason": reason})
            preview.append({"wizard_id": self.id, "row_number": row_number, "status": status, "reason": reason, "data_json": json.dumps(row, default=str)})
        self.env["fedex.import.preview"].create(preview)
        if errors:
            output = io.StringIO()
            writer = csv.DictWriter(output, fieldnames=["row", "reason"])
            writer.writeheader(); writer.writerows(errors)
            self.error_report = base64.b64encode(output.getvalue().encode())
            self.error_report_filename = "fedex_import_errors.csv"
        self.state = "preview"
        return {"type": "ir.actions.act_window", "res_model": self._name, "res_id": self.id, "view_mode": "form", "target": "new"}

    def action_import(self):
        self.ensure_one()
        rows = self._rows()
        model = self.env[ENTITY_MODELS[self.entity]]
        # One transaction per entity: the method runs as one entity transaction in Odoo's request transaction.
        for index, row in enumerate(rows, 2):
            duplicate = self._find_duplicate(row)
            preview_match = self.preview_line_ids.filtered(lambda x: x.row_number == index)
            if preview_match and preview_match.status == "error":
                continue
            vals = self._build_vals(row)
            if duplicate:
                if self.import_mode == "skip":
                    continue
                duplicate.write(vals)
            else:
                model.create(vals)
        return {"type": "ir.actions.act_window_close"}

    def action_download_template(self):
        self.ensure_one()
        return {"type": "ir.actions.act_url", "url": "/fedex_repository/import/template/%s" % self.entity, "target": "self"}


class FedexImportPreview(models.TransientModel):
    _name = "fedex.import.preview"
    _description = "Import Preview"
    _order = "row_number"

    wizard_id = fields.Many2one("fedex.import.wizard", required=True, ondelete="cascade")
    row_number = fields.Integer(required=True)
    status = fields.Selection([("new", "New"), ("update", "Update"), ("duplicate", "Duplicate"), ("error", "Error")], required=True)
    reason = fields.Char()
    data_json = fields.Text()
